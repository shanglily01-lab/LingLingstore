# -*- coding: utf-8 -*-
"""Generate Brazil cross-border trade inquiry Word document."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

DOC_DIR = os.path.dirname(__file__)
OUTPUT_CN = os.path.join(DOC_DIR, "巴西外贸-清关海外仓询价清单.docx")
OUTPUT_EN = os.path.join(DOC_DIR, "brazil-trade-inquiry-checklist.docx")


def set_cell_shading(cell, fill: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return h


def add_para(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    return p


def add_bullet(doc, text, size=11):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    return p


def add_numbered(doc, text, size=11):
    p = doc.add_paragraph(style="List Number")
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    return p


def add_quote_box(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1)
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_shading(hdr[i], "D9E2F3")
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.name = "Microsoft YaHei"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                run.font.size = Pt(10)
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
            for p in cells[ci].paragraphs:
                for run in p.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                    run.font.size = Pt(10)
    doc.add_paragraph()
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("巴西跨境外贸实操手册\n（清关 · 海外仓 · 销售 · 询价清单）")
    tr.bold = True
    tr.font.size = Pt(18)
    tr.font.name = "Microsoft YaHei"
    tr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("适用场景：中国海运 → 圣保罗海外仓 → 亚马逊/美客多自发货 | 品类：纸质扑克牌")
    sr.font.size = Pt(11)
    sr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    sr.font.name = "Microsoft YaHei"
    sr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    doc.add_paragraph()

    # Section 0: Core summary
    add_heading(doc, "一、核心要点速记（外贸小白版）", 1)
    add_para(doc, "全流程一句话：", bold=True)
    add_para(doc, "货能不能卖，看清关能不能过；能不能按时卖，看海外仓能不能接好；能不能赚钱，看定价和运营。")

    add_para(doc, "四大服务商分工：", bold=True)
    add_table(doc, ["角色", "负责什么", "什么时候找"], [
        ["中国货代", "订舱、出口报关、海运", "第一批货发出前"],
        ["巴西清关行", "进口申报、缴税、海关放行", "与货代同时找，需互相配合"],
        ["圣保罗海外仓", "收货、存货、按单拣货发货", "清关前就要签好合同"],
        ["亚马逊/美客多", "卖货、收款、生成快递面单", "入仓前 listing 先备好"],
    ])

    add_para(doc, "首单时间线参考：", bold=True)
    add_table(doc, ["阶段", "周期"], [
        ["CNPJ + 清关行/海外仓对接", "2–8 周"],
        ["采购 + 海运（中国→桑托斯）", "35–45 天"],
        ["巴西清关", "7–15 天（查验另计）"],
        ["入海外仓 + 可售上架", "3–7 天"],
        ["首单全程合计", "约 3–4 个月"],
    ])

    # Section: Customs
    add_heading(doc, "二、巴西海关清关 — 关键要点", 1)
    add_para(doc, "清关 = 货船到桑托斯港后，向巴西海关申报、缴税、获准进入巴西。放行前货不能进海外仓，也不能销售。")

    add_para(doc, "必备条件：", bold=True)
    add_bullet(doc, "CNPJ（巴西公司税号）+ RADAR（进口资质）")
    add_bullet(doc, "巴西清关行（实际操办进口申报）")
    add_bullet(doc, "NCM 税号：扑克牌一般为 9504.40.00（清关行最终确认）")
    add_bullet(doc, "完整单证：Commercial Invoice、Packing List、提单 B/L，信息必须与实物一致")

    add_para(doc, "清关流程：", bold=True)
    add_quote_box(doc, [
        "船到桑托斯港 → 清关行提交进口申报 → 海关审核（可能查验）",
        "→ 缴纳进口税（II + IPI + PIS/COFINS + ICMS）→ 放行 → 拖车至圣保罗海外仓",
    ])

    add_para(doc, "常见翻车点：", bold=True)
    add_table(doc, ["风险", "后果", "预防"], [
        ["NCM 税号错误", "补税、扣货、延误", "发货前让清关行书面确认"],
        ["发票与装箱单不符", "海关不放行", "数量、SKU、金额三方统一"],
        ["无 CNPJ/RADAR", "无法清关", "开店/发货前先搞定"],
        ["税算不足", "卖一单亏一单", "清关行做含税到岸价模拟后再采购"],
    ])

    add_heading(doc, "三、海外仓收货 — 关键要点", 1)
    add_para(doc, "清关放行后，货从港口拖到圣保罗 3PL 仓库，完成卸柜、清点、上架后，系统出现「可售库存」，才能开始销售。")

    add_para(doc, "入仓前必须发给仓库（ASN 预报）：", bold=True)
    add_bullet(doc, "柜号 / 提单号、预计到港时间 ETA")
    add_bullet(doc, "SKU 明细表（每款多少箱、每箱多少副）")
    add_bullet(doc, "EAN 条码、外箱箱唛样式")
    add_bullet(doc, "是否需要代贴 FNSKU / EAN 标签")

    add_para(doc, "常见翻车点：", bold=True)
    add_table(doc, ["风险", "后果", "预防"], [
        ["箱唛未贴或贴错", "仓库无法分拣", "中国出货前按仓库规范贴好"],
        ["未做入仓预报", "延误上架", "开船前发送 ASN"],
        ["平台已开售但仓内无货", "超卖、罚店", "入仓完成后再改可售库存"],
        ["两平台库存不同步", "超卖", "共仓统一扣减，设 5–10% 安全库存"],
    ])

    add_heading(doc, "四、双平台销售 — 关键要点", 1)
    add_para(doc, "模式：亚马逊 FBM + 美客多 Mercado Envios 自发货（非 FBA/非 Full），共用圣保罗海外仓库存。")

    add_quote_box(doc, [
        "买家下单 → 平台通知你 → 推订单给海外仓（含面单）",
        "→ 仓库拣货贴单 → 交本地快递 → 回填运单号 → 买家签收 → 平台结算",
    ])

    add_para(doc, "定价底线公式：", bold=True)
    add_quote_box(doc, [
        "售价 ≥ 到仓成本 + 仓租 + 出库费 + 本地快递费 + 平台佣金(12–16%)",
        "       + 广告(5–15%) + 退货损耗(3–8%) + 目标利润(≥20%)",
    ])

    add_para(doc, "扑克牌巴西市场参考零售价（2026年初调研）：", bold=True)
    add_table(doc, ["品类", "主流价格带 (BRL)", "约合 (CNY)"], [
        ["普通纸扑克（单副）", "R$ 18 – 25", "¥ 24 – 33"],
        ["普通纸扑克（双副装）", "R$ 35 – 45", "¥ 46 – 59"],
        ["德州 PVC 扑克（单副）", "R$ 50 – 72（Copag 标杆）", "¥ 65 – 94"],
        ["进口/generic 塑料扑克", "R$ 20 – 35", "¥ 26 – 46"],
        ["建议你的零售价（纸牌单副）", "R$ 22 – 32", "¥ 29 – 42"],
    ])

    # Inquiry sections
    add_heading(doc, "五、发给巴西清关行的询价问题（10 题）", 1)
    add_para(doc, "背景说明（可复制发给对方）：", bold=True)
    add_quote_box(doc, [
        "我司计划从中国海运扑克牌（NCM 预计 9504.40.00）至桑托斯港，",
        "清关后拖至圣保罗海外仓。我司有 CNPJ 及 RADAR 进口资质。",
        "请逐项回复以下问题，并提供书面报价。",
    ])

    customs_q = [
        "扑克牌（纸质/塑料，NCM 9504.40.00）能否正常进口？是否需要 INMETRO 或其他认证？",
        "请按 CIF 价 USD ______ / 数量 ______ 副，做一份含税到岸价模拟（II、IPI、PIS/COFINS、ICMS 分项列出）。",
        "清关正常需要几天？若海关查验，通常额外多少天？",
        "清关费、港杂、换单、拖柜至圣保罗仓各是多少？能否报「港到仓」一条龙总价？",
        "我需要提供哪些单证？Invoice、Packing List 有无固定模板？",
        "进口申报用我司 CNPJ，能否开具 Nota Fiscal（巴西发票）？如何配合做账？",
        "若 NCM 归类或品名有误导致补税/扣货，如何处理？责任如何划分？",
        "是否支持 LCL 拼柜首单？最小接单量是多少？",
        "能否与我指定的圣保罗海外仓对接？清关放行后直接拖柜到仓是否可行？",
        "全程对接联系人是谁？能否与货代、海外仓建立三方沟通群跟进首单？",
    ]
    for i, q in enumerate(customs_q, 1):
        add_numbered(doc, q)

    add_para(doc, "请对方按以下格式回复：", bold=True)
    add_quote_box(doc, [
        "到岸成本明细：",
        "清关周期：",
        "所需单证：",
        "总报价（含/不含拖柜）：",
        "注意事项：",
    ])

    add_heading(doc, "六、发给圣保罗海外仓的询价问题（12 题）", 1)
    add_para(doc, "背景说明（可复制发给对方）：", bold=True)
    add_quote_box(doc, [
        "货物从桑托斯港清关后入贵仓，库存供亚马逊 FBM 与美客多 Mercado Envios",
        "自发货共用。品类为扑克牌（纸质）。请逐项回复并提供报价。",
    ])

    wh_q = [
        "仓库具体地址？距桑托斯港多远？是否常做桑托斯拖柜入仓？",
        "能否接收 40HQ / LCL 拼柜？卸柜费、入库上架费如何计算（按柜/按箱/按件）？",
        "仓储费如何收取（每 m³/托盘/天）？有无最低消费？",
        "是否支持亚马逊 FBM + 美客多 Mercado Envios 双平台出库？",
        "能否打印 Mercado Envios 面单和亚马逊面单？打印机/格式是否兼容？",
        "能否代贴 EAN 条码、FNSKU？贴标费多少？",
        "订单如何推送？（API / 邮件 / Excel）有无模板？能否方便回传运单号？",
        "多平台共用一个库存，能否防止超卖？系统如何同步？",
        "是否接收买家退件？退件检测、重新上架如何收费？",
        "入仓需要哪些资料？（ASN 预报、箱唛、SKU 明细表）有无模板？",
        "从收货到可售库存上架，通常需要几天？",
        "首单能否安排测试出库（亚马逊、美客多各 1 单）？如何收费？",
    ]
    for i, q in enumerate(wh_q, 1):
        add_numbered(doc, q)

    add_para(doc, "请对方按以下格式回复：", bold=True)
    add_quote_box(doc, [
        "卸柜/上架/仓储/出库单价：",
        "是否支持双平台：",
        "推单方式：",
        "退件方案：",
        "入仓所需文件：",
    ])

    add_heading(doc, "七、发给中国货代的询价问题（8 题）", 1)
    add_para(doc, "背景说明（可复制发给对方）：", bold=True)
    add_quote_box(doc, [
        "需从 __________ 港口海运扑克牌至巴西桑托斯，配合巴西清关行与圣保罗海外仓。",
        "请逐项回复并提供报价。",
    ])

    ff_q = [
        "从 __________ 到桑托斯，LCL 拼柜当前海运费多少？（每 CBM 或每吨）",
        "中国段出口报关费用？是否支持买单出口？",
        "正常航程多少天？最近是否有绕航或延误？",
        "能否与我的巴西清关行直接对接单证？我需要提供哪些文件？",
        "Packing List、Invoice 有无标准模板？是否协助审核？",
        "能否协调「港到仓」全程（海运+巴西清关+拖柜）？还是只服务到港？",
        "首单 LCL 最小起运量？扑克牌建议多大体积试单？",
        "若船期延误或甩柜，如何通知？有无预案？",
    ]
    for i, q in enumerate(ff_q, 1):
        add_numbered(doc, q)

    add_heading(doc, "八、平台入驻确认问题（可选）", 1)
    add_para(doc, "亚马逊巴西：", bold=True)
    add_bullet(doc, "用 CNPJ 入驻 FBM 当前要求和审核周期？")
    add_bullet(doc, "FBM 发货地设圣保罗有无限制？")
    add_bullet(doc, "Handling Time 建议设多少天？")

    add_para(doc, "美客多巴西：", bold=True)
    add_bullet(doc, "自发货（非 Full）如何绑定 Mercado Envios？")
    add_bullet(doc, "面单是否必须由卖家推给仓库打印？")
    add_bullet(doc, "与亚马逊共用同一海外仓库存有无平台限制？")

    add_heading(doc, "九、报价对比表（收到回复后填写）", 1)
    add_table(doc, ["对比项", "清关行 A", "清关行 B", "海外仓 A", "海外仓 B"], [
        ["含税到岸价（每副）", "", "", "", ""],
        ["港到仓总周期", "", "", "", ""],
        ["是否支持 LCL 首单", "", "", "", ""],
        ["能否双平台出库", "", "", "", ""],
        ["沟通响应速度", "", "", "", ""],
        ["综合评分（1-5）", "", "", "", ""],
    ])

    add_para(doc, "决策原则：", bold=True)
    add_bullet(doc, "清关行：优先选报价透明 + 有扑克牌/NCM 经验 + 能对接你的海外仓")
    add_bullet(doc, "海外仓：优先选双平台出库成熟 + 退件方案清楚 + 桑托斯拖柜方便")
    add_bullet(doc, "货代：优先选能三方协同，不必单纯追求最低价")

    add_heading(doc, "十、首单 SKU 入仓映射表（模板）", 1)
    add_table(doc, [
        "SKU 名称", "EAN 条码", "亚马逊 ASIN", "美客多 Item ID",
        "每箱副数", "箱数", "总副数", "箱规(cm)", "单箱毛重(kg)", "备注",
    ], [
        ["扑克-经典款", "", "", "", "", "", "", "", "", ""],
        ["扑克-主题款", "", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "", "", "", "", "", ""],
    ])

    add_heading(doc, "十一、首单行动 Checklist", 1)
    checklist = [
        "[ ] 注册或搞定 CNPJ + RADAR",
        "[ ] 找清关行：做 1 万副含税到岸价模拟",
        "[ ] 找圣保罗海外仓：确认双平台出库 + 桑托斯拖柜报价",
        "[ ] 找中国货代：LCL 海运 + 出口报关报价",
        "[ ] 清关行、货代、海外仓拉三方群",
        "[ ] 确定 2 个 SKU 打样（含葡语盒标 + EAN）",
        "[ ] 注册亚马逊 BR + 美客多卖家账号",
        "[ ] 葡语 listing 备好（库存先设 0）",
        "[ ] 中国出货：贴箱唛、准备 Invoice + Packing List",
        "[ ] 开船前：发 ASN 入仓预报给海外仓",
        "[ ] 货到仓：核对实收数量 = 单证数量",
        "[ ] 可售库存确认后：平台改库存、开始销售",
        "[ ] 首单全程跟单：来单→推仓→出库→填运单号→签收",
    ]
    for item in checklist:
        add_para(doc, item, size=11)

    add_heading(doc, "十二、首单每周跟单话术", 1)
    add_table(doc, ["问谁", "说什么"], [
        ["中国货代", "货现在在哪？预计哪天到桑托斯？"],
        ["巴西清关行", "申报提交了吗？有无查验？预计哪天放行？"],
        ["圣保罗海外仓", "柜到了吗？上架完成多少？当前可售库存是多少？"],
    ])

    add_heading(doc, "十三、给清关行的询价邮件/消息模板", 1)
    add_quote_box(doc, [
        "Olá / 您好，",
        "",
        "我司计划从中国海运扑克牌至巴西，详情如下：",
        "  · 品类：纸质扑克牌（Baralho / Playing Cards）",
        "  · NCM 预计：9504.40.00",
        "  · 目的港：Santos",
        "  · 最终目的地：圣保罗海外仓（地址：__________）",
        "  · 运输方式：LCL 拼柜 / FCL 整柜（首单）",
        "  · 预计数量：__________ 副",
        "  · 预计 CIF 货值：USD __________",
        "  · 进口商：我司 CNPJ __________",
        "",
        "请协助：",
        "  1. 确认 NCM 及是否需要 INMETRO 等认证",
        "  2. 提供含税到岸价模拟（分项列税）",
        "  3. 报清关费 + 港杂 + 拖柜至圣保罗仓的总价",
        "  4. 告知所需单证及预计清关周期",
        "",
        "附件：产品图片、材质说明（如有）",
        "期待回复，谢谢！",
    ])

    # Footer note
    doc.add_paragraph()
    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = foot.add_run("— 文档生成日期：2026 年 6 月 | 仅供内部实操参考 —")
    fr.font.size = Pt(9)
    fr.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    fr.font.name = "Microsoft YaHei"
    fr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    doc.save(OUTPUT_CN)
    doc.save(OUTPUT_EN)
    print(f"Saved: {OUTPUT_CN}")
    print(f"Saved: {OUTPUT_EN}")


if __name__ == "__main__":
    build()
