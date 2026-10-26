# -*- coding: utf-8 -*-
"""Generate friend inquiry Word doc: WeChat short + full question list."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

DOC_DIR = os.path.dirname(__file__)
OUTPUT = os.path.join(DOC_DIR, "巴西外贸-问朋友常见问题.docx")

INTRO = (
    "我在做中国纸扑克牌进巴西，计划桑托斯清关、圣保罗海外仓，"
    "亚马逊 FBM + 美客多自发货，NCM 9504.40.00，首单 LCL 试跑。"
    "外贸新手，最担心清关、入仓和销售三块，想请教你/有没有资源推荐。"
)

WECHAT_SHORT = [
    "1. 纸扑克牌 NCM 9504.40.00 对不对？要不要 INMETRO？",
    "2. 1 万副 LCL，桑托斯清关+进圣保罗仓，含税到岸成本大概多少？税占 CIF 几成？",
    "3. 没有 CNPJ/RADAR 能不能做？注册大概多久、多少钱？",
    "4. 清关行、货代、圣保罗海外仓，你有靠谱推荐吗？要不要三方提前拉群？",
    "5. 清关正常几天？被海关查验一般拖多久？放行前会不会产生港口堆存费？",
    "6. 圣保罗仓能不能同时做亚马逊 FBM + 美客多自发货？面单各试打一张行不行？",
    "7. 双平台共用一个仓，怎么防超卖？来单怎么推仓、截单几点？",
    "8. 亚马逊和美客多扑克牌大概卖多少钱一副？你更建议先做哪个平台？",
    "9. 首单试跑（5千–2万副）从头到尾要准备多少流动资金？",
    "10. 你当时踩过最大的坑是什么？如果重来会怎么做？",
]

SECTIONS = [
    ("一、整体路线", [
        "我这种货（纸扑克牌，首单 5,000–20,000 副，LCL 拼柜），走桑托斯 → 圣保罗海外仓这条路靠谱吗？",
        "货代、清关行、海外仓是自己分别找，还是有人能打包推荐？要不要提前拉群？",
        "从中国发货到圣保罗仓可卖，正常要多久？哪个环节最容易拖时间？",
        "首单预算大概要留多少？进口税和杂费一般占货值多少？",
        "没有巴西公司（CNPJ）能不能做？注册 CNPJ + RADAR 大概要多久、多少钱？",
    ]),
    ("二、清关", [
        "扑克牌 NCM 9504.40.00 对不对？要不要 INMETRO 或其他认证？",
        "帮我按 1 万副估一下含税到岸成本：II、IPI、PIS/COFINS、ICMS 各多少？",
        "清关是用我自己的 CNPJ 还是可以用代理？各有什么风险？",
        "船到桑托斯后，清关正常要几天？被海关查验一般要拖多久？",
        "清关需要提前准备哪些单证？（Invoice、Packing List、提单等）",
        "货值、品名、箱规和实物有出入，会不会扣货或补税？",
        "进口税是清关行代缴还是我直接付？大概什么时候要准备好钱？",
        "清关放行之前货能不能先挪到仓？会不会产生港口堆存费？",
        "你认识的清关行桑托斯 → 圣保罗这条线熟不熟？能不能港到仓一条龙？",
    ]),
    ("三、海外仓", [
        "圣保罗有没有你用过、靠谱的 3PL？能不能接桑托斯拖柜的 LCL 货？",
        "卸柜、清点、上架一般要几天？按什么收费？",
        "仓库能不能同时帮亚马逊 FBM + 美客多自发货打面单、贴单、出库？",
        "亚马逊和美客多面单格式仓库能不能打？首单前能不能各试打一张？",
        "来单后怎么推给仓库？截单时间是几点？",
        "实收数量和装箱单不符（少货、错 SKU），一般怎么处理？",
        "能不能代贴 EAN / FNSKU？贴标怎么收费？",
        "退件怎么处理？退到哪里、能否重新上架、运费谁出？",
        "双平台共用一个仓，怎么防超卖？",
    ]),
    ("四、中国货代 / 海运", [
        "从中国哪个港出到桑托斯，LCL 海运费大概多少？周期多长？",
        "中国出口报关货代包吗？我需要提供什么资料？",
        "货代能不能和巴西清关行、圣保罗仓提前对接？",
        "首单 LCL，箱唛、发票、箱单有什么格式要求？",
        "海运途中怎么跟踪？开船、到港一般谁通知我？",
    ]),
    ("五、亚马逊 / 美客多销售", [
        "扑克牌在巴西，亚马逊和美客多哪个更好卖？更建议先做哪个？",
        "FBM 从圣保罗发，Handling Time 设几天比较合理？",
        "美客多是自发货还是 Full？我这种共仓模式哪种更合适？",
        "扑克牌在巴西大概卖多少钱一副？（德州扑克 vs 普通扑克有区别吗？）",
        "平台佣金 + 广告 + 仓配加起来，大概占售价多少？",
        "Listing 葡语怎么搞？有没有靠谱的人或工具？",
        "收款用 Payoneer 还是别的？回款周期大概多久？",
        "上架前库存要不要先设 0，等入仓完成再开卖？",
    ]),
    ("六、产品 & 合规", [
        "包装文案写「扑克牌 / 纸牌」就行，还是要避免某些词（如儿童玩具）？",
        "每副牌需要 EAN 条码吗？亚马逊和美客多各要什么格式？",
        "外箱箱唛要贴哪些信息？葡语标签有没有强制要求？",
        "工厂出货前，质检和贴标有什么必查项？",
    ]),
    ("七、钱和时间", [
        "首单试跑，大概要准备多少流动资金？",
        "哪个时间点必须打款？（定金、尾款、税金、仓费）",
        "首单小批量试跑（5,000 副），量会不会太小、运费不划算？",
        "你当时首单踩过最大的坑是什么？如果重来会怎么做？",
    ]),
]


def set_cell_shading(cell, fill: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def style_run(run, size=11, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        style_run(run, size=16 if level == 1 else 13, bold=True)
    return h


def add_para(doc, text, bold=False, size=11, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(text)
    style_run(run, size=size, bold=bold)
    return p


def add_bullet(doc, text, size=11):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    style_run(run, size=size)
    return p


def add_quote_box(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.right_indent = Cm(0.5)
    run = p.add_run(text)
    style_run(run, size=10, color=RGBColor(0x33, 0x33, 0x33))
    return p


def build():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("巴西外贸 · 问朋友常见问题清单")
    style_run(run, size=18, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("纸扑克牌 · 桑托斯清关 · 圣保罗海外仓 · 亚马逊/美客多")
    style_run(run, size=11, color=RGBColor(0x66, 0x66, 0x66))

    doc.add_paragraph()

    add_heading(doc, "背景说明（可复制发给朋友）", level=2)
    add_quote_box(doc, INTRO)

    doc.add_paragraph()
    add_heading(doc, "微信短版 · 10 个核心问题", level=1)
    add_para(
        doc,
        "适合直接复制到微信/WhatsApp，一次问完最关键的事。",
        size=10,
    )
    doc.add_paragraph()
    for q in WECHAT_SHORT:
        add_para(doc, q, size=11)

    doc.add_page_break()

    add_heading(doc, "完整版 · 分环节问题清单", level=1)
    add_para(doc, "和朋友语音或见面时，可以按章节逐项聊。", size=10)
    doc.add_paragraph()

    for section_title, questions in SECTIONS:
        add_heading(doc, section_title, level=2)
        for i, q in enumerate(questions, 1):
            add_numbered = doc.add_paragraph(style="List Number")
            run = add_numbered.add_run(q)
            style_run(run, size=11)
        doc.add_paragraph()

    add_heading(doc, "备注栏（问完后填写）", level=2)
    add_bullet(doc, "清关行推荐：________________  联系方式：________________")
    add_bullet(doc, "海外仓推荐：________________  联系方式：________________")
    add_bullet(doc, "货代推荐：________________    联系方式：________________")
    add_bullet(doc, "含税到岸成本估算：________________")
    add_bullet(doc, "首单建议采购量：________________")
    add_bullet(doc, "其他提醒：________________")

    doc.save(OUTPUT)
    print(f"Generated: {OUTPUT}")


if __name__ == "__main__":
    build()
